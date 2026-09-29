import io
import pytest
import docx
from typing import Dict, Any
from app.schemas.ocr import (
    DutySlipExtractedFields,
    ExtractedField,
    CalculatedValidationSummary
)
from app.schemas.company import RateConfig, CompanyCreate, CompanyResponse
from app.services.ocr.validator import normalize_and_validate_extraction
from app.services.word.generator import WordBillGenerator


def test_schema_backward_compatibility_missing_optional_fields():
    """
    Test that legacy extraction records without Step 9.1 fields (passenger_name, booked_by,
    parking, service_type, vehicle_category) load cleanly without schema errors.
    """
    legacy_data = {
        "duty_slip_no": {"value": "101", "confidence": 0.95},
        "date": {"value": "01-09-2026", "confidence": 0.95},
        "driver_name": {"value": "Legacy Driver", "confidence": 0.92},
        "vehicle_number": {"value": "TS09AA1111 Sedan", "confidence": 0.94},
        "meter_start": {"value": 1000, "confidence": 0.96},
        "meter_return": {"value": 1080, "confidence": 0.96},
        "total_km": {"value": 80, "confidence": 0.97},
        "starting_time": {"value": "09:00 AM", "confidence": 0.93},
        "closing_time": {"value": "05:00 PM", "confidence": 0.93},
        "extra_km": {"value": 0, "confidence": 0.90},
        "extra_hours": {"value": 0.0, "confidence": 0.90},
        "reporting_place": {"value": "Office", "confidence": 0.90},
        "tour_location": {"value": "Local", "confidence": 0.90},
        "party_name": {"value": "Legacy Client", "confidence": 0.90},
        "customer_name": {"value": "Legacy Client", "confidence": 0.90},
        "base_package": {"value": "8/80 = 2500", "confidence": 0.95},
        "extra_km_charge": {"value": "0", "confidence": 0.95},
        "extra_hour_charge": {"value": "0", "confidence": 0.95},
        "bata": {"value": "0", "confidence": 0.95},
        "toll": {"value": "0", "confidence": 0.95},
        "total_amount": {"value": "2500", "confidence": 0.95}
    }

    # Should instantiate without raising ValidationError
    fields_obj = DutySlipExtractedFields(**legacy_data)
    assert fields_obj.duty_slip_no.value == "101"
    # New fields should default to value None
    assert fields_obj.passenger_name.value is None
    assert fields_obj.booked_by.value is None
    assert fields_obj.parking.value is None
    assert fields_obj.service_type.value is None
    assert fields_obj.vehicle_category.value is None

    # Normalization & validation must succeed without crashing on legacy inputs
    norm_fields, val_summary, conf, needs_rev, reasons = normalize_and_validate_extraction(
        raw_fields=legacy_data,
        stored_duty_slip_no="101",
        has_back_scan=True
    )
    assert norm_fields.duty_slip_no.value == "101"
    assert norm_fields.passenger_name.value is None
    assert val_summary.calculated_parking_amount is None
    assert val_summary.is_math_consistent is True


# =========================================================================
# FOCUSED SEMANTIC TESTS: TEST 1 THROUGH TEST 6 (STEP 9.1 CORRECTIVE PATCH)
# =========================================================================

def test_1_customer_name_not_equal_passenger_name_when_passenger_absent():
    """
    TEST 1:
    customer_name != passenger_name.
    When customer_name is provided and passenger_name is None, passenger_name remains None.
    customer_name / party_name must NOT be automatically mapped to passenger_name.
    """
    raw_data = {
        "duty_slip_no": "201",
        "customer_name": "Windsor Machines Limited",
        "passenger_name": None
    }
    fields_obj, _, _, _, _ = normalize_and_validate_extraction(
        raw_fields=raw_data,
        stored_duty_slip_no="201",
        has_back_scan=False
    )
    assert fields_obj.customer_name.value == "Windsor Machines Limited"
    assert fields_obj.passenger_name.value is None


def test_2_customer_name_company_and_passenger_name_individual_independent():
    """
    TEST 2:
    customer_name can be a company while passenger_name is an individual.
    Example:
      customer_name = "Windsor Machines Limited"
      passenger_name = "Mr. Vipul"
    Both remain independent.
    """
    raw_data = {
        "duty_slip_no": "202",
        "customer_name": "Windsor Machines Limited",
        "passenger_name": "Mr. Vipul"
    }
    fields_obj, _, _, _, _ = normalize_and_validate_extraction(
        raw_fields=raw_data,
        stored_duty_slip_no="202",
        has_back_scan=False
    )
    assert fields_obj.customer_name.value == "Windsor Machines Limited"
    assert fields_obj.passenger_name.value == "Mr. Vipul"
    assert fields_obj.customer_name.value != fields_obj.passenger_name.value


def test_3_booked_by_independent_from_customer_and_driver():
    """
    TEST 3:
    booked_by remains independent from both customer_name and driver_name.
    """
    raw_data = {
        "duty_slip_no": "203",
        "customer_name": "Welspun Corp",
        "driver_name": "Driver Govind",
        "booked_by": "Welspun Officer Sharma"
    }
    fields_obj, _, _, _, _ = normalize_and_validate_extraction(
        raw_fields=raw_data,
        stored_duty_slip_no="203",
        has_back_scan=False
    )
    assert fields_obj.customer_name.value == "Welspun Corp"
    assert fields_obj.driver_name.value == "Driver Govind"
    assert fields_obj.booked_by.value == "Welspun Officer Sharma"
    assert fields_obj.booked_by.value != fields_obj.driver_name.value
    assert fields_obj.booked_by.value != fields_obj.customer_name.value

    # When booked_by is absent, it must NOT be populated from customer_name or driver_name
    raw_data_no_booked = {
        "duty_slip_no": "203B",
        "customer_name": "Welspun Corp",
        "driver_name": "Driver Govind",
        "booked_by": None
    }
    fields_obj2, _, _, _, _ = normalize_and_validate_extraction(
        raw_fields=raw_data_no_booked,
        stored_duty_slip_no="203B",
        has_back_scan=False
    )
    assert fields_obj2.booked_by.value is None


def test_4_driver_name_never_used_as_booked_by():
    """
    TEST 4:
    driver_name is never used as booked_by in Word document generation.
    """
    generator = WordBillGenerator()
    duty_slip = {
        "duty_slip_no": "204",
        "driver_name": "Driver Ramesh",
        "passenger_name": "Mr. Vipul",
        "booked_by": None,
        "company_name": "Windsor Machines Limited",
        "base_package": "2500",
        "total_amount": "2500"
    }
    docx_bytes = generator.generate_single_bill_docx_bytes(
        duty_slip_no="204",
        fields_dict=duty_slip,
        bill_no="Bill 01",
        company_name="Windsor Machines Limited"
    )
    doc = docx.Document(io.BytesIO(docx_bytes))
    paragraphs_text = [p.text for p in doc.paragraphs]
    booked_by_p = [p for p in paragraphs_text if "Booked by" in p]
    assert len(booked_by_p) > 0
    assert "Booked by : Windsor Machines Limited" in booked_by_p[0]
    assert "Driver Ramesh" not in booked_by_p[0]


def test_5_generator_does_not_put_customer_name_into_for_field_when_passenger_absent():
    """
    TEST 5:
    Generator does not put customer_name into the "For :" field merely because passenger_name is absent.
    """
    generator = WordBillGenerator()
    duty_slip = {
        "duty_slip_no": "205",
        "customer_name": "Windsor Machines Limited",
        "passenger_name": None,
        "driver_name": "Driver Ramesh",
        "base_package": "2500",
        "total_amount": "2500"
    }
    docx_bytes = generator.generate_single_bill_docx_bytes(
        duty_slip_no="205",
        fields_dict=duty_slip,
        bill_no="Bill 01",
        company_name="Windsor Machines Limited"
    )
    doc = docx.Document(io.BytesIO(docx_bytes))
    paragraphs_text = [p.text for p in doc.paragraphs]

    for_p = [p for p in paragraphs_text if p.startswith("For :")]
    assert len(for_p) > 0
    # Must be "For :" without the company/customer name
    assert for_p[0].strip() == "For :"
    assert "Windsor Machines Limited" not in for_p[0]

    # But customer_name MUST appear under To, in the recipient block
    to_p_index = next(i for i, p in enumerate(paragraphs_text) if p.startswith("To, "))
    recipient_p = paragraphs_text[to_p_index + 1]
    assert "Windsor Machines Limited" in recipient_p


def test_6_existing_customer_name_functionality_continues_to_work():
    """
    TEST 6:
    Existing customer_name-based functionality continues to work.
    Recipient block displays customer_name, and passenger displays passenger_name.
    """
    generator = WordBillGenerator()
    duty_slip = {
        "duty_slip_no": "206",
        "customer_name": "Acme Global Corporation",
        "passenger_name": "Dr. Smith",
        "driver_name": "Driver Ramesh",
        "base_package": "2500",
        "total_amount": "2500"
    }
    docx_bytes = generator.generate_single_bill_docx_bytes(
        duty_slip_no="206",
        fields_dict=duty_slip,
        bill_no="Bill 01",
        company_name="Acme Global Corporation"
    )
    doc = docx.Document(io.BytesIO(docx_bytes))
    paragraphs_text = [p.text for p in doc.paragraphs]

    to_p_index = next(i for i, p in enumerate(paragraphs_text) if p.startswith("To, "))
    recipient_p = paragraphs_text[to_p_index + 1]
    assert "Acme Global Corporation" in recipient_p

    for_p = [p for p in paragraphs_text if p.startswith("For :")]
    assert len(for_p) > 0
    assert for_p[0].strip() == "For :Dr. Smith"


# =========================================================================
# ADDITIONAL STEP 9.1 FOUNDATION TESTS
# =========================================================================

def test_parking_field_behavior_absent_and_present():
    """
    Test that parking charge can be either absent (optional) or present.
    When present, it must be factored into mathematical validation.
    """
    # 1. Parking Absent
    raw_absent = {
        "duty_slip_no": "301",
        "base_package": "2500",
        "extra_km_charge": "1005",
        "extra_hour_charge": "225",
        "bata": "250",
        "toll": "40",
        "total_amount": "4020"
    }
    fields_absent, summary_absent, _, needs_rev_absent, _ = normalize_and_validate_extraction(
        raw_fields=raw_absent,
        stored_duty_slip_no="301",
        has_back_scan=True
    )
    assert fields_absent.parking.value is None
    assert summary_absent.calculated_parking_amount is None
    assert summary_absent.is_math_consistent is True
    # Parking missing should NOT cause review
    assert fields_absent.parking.needs_review is False

    # 2. Parking Present & Math Consistent
    # 2500 + 1005 + 225 + 250 + 40 + 150 (parking) = 4170
    raw_present = {
        "duty_slip_no": "302",
        "base_package": "2500",
        "extra_km_charge": "1005",
        "extra_hour_charge": "225",
        "bata": "250",
        "toll": "40",
        "parking": "150",
        "total_amount": "4170"
    }
    fields_present, summary_present, _, _, _ = normalize_and_validate_extraction(
        raw_fields=raw_present,
        stored_duty_slip_no="302",
        has_back_scan=True
    )
    assert fields_present.parking.value == "150"
    assert summary_present.calculated_parking_amount == 150.0
    assert summary_present.calculated_total_amount == 4170.0
    assert summary_present.is_math_consistent is True

    # 3. Parking Present but Total Mismatch
    raw_mismatch = {
        "duty_slip_no": "303",
        "base_package": "2500",
        "extra_km_charge": "1005",
        "extra_hour_charge": "225",
        "bata": "250",
        "toll": "40",
        "parking": "150",
        "total_amount": "4020"  # Forgot to add parking into total
    }
    fields_mismatch, summary_mismatch, _, needs_rev_mismatch, reasons = normalize_and_validate_extraction(
        raw_fields=raw_mismatch,
        stored_duty_slip_no="303",
        has_back_scan=True
    )
    assert summary_mismatch.is_math_consistent is False
    assert fields_mismatch.total_amount.needs_review is True
    assert "validation_failure" in reasons


def test_company_rate_config_representation():
    """
    Test that RateConfig can accurately represent observed company and vehicle tariffs:
    - Windsor Sedan (2500 base, 15 extra km, 150 extra hr)
    - Welspun Sedan (2100 base, 14 extra km, 120 extra hr)
    - Crysta (3750 base, 20 extra km, 300 extra hr)
    - Outstation rate model (300 km min, 15/km, 300 driver bata)
    """
    windsor_sedan = RateConfig(
        vehicle_category="sedan",
        service_type="local",
        base_package_rate=2500.0,
        base_km=80.0,
        base_hours=8.0,
        extra_km_rate=15.0,
        extra_hour_rate=150.0,
        driver_bata_rate=None
    )

    welspun_sedan = RateConfig(
        vehicle_category="sedan",
        service_type="local",
        base_package_rate=2100.0,
        base_km=80.0,
        base_hours=8.0,
        extra_km_rate=14.0,
        extra_hour_rate=120.0,
        driver_bata_rate=None
    )

    crysta_local = RateConfig(
        vehicle_category="crysta",
        service_type="local",
        base_package_rate=3750.0,
        base_km=80.0,
        base_hours=8.0,
        extra_km_rate=20.0,
        extra_hour_rate=300.0,
        driver_bata_rate=None
    )

    outstation_sedan = RateConfig(
        vehicle_category="sedan",
        service_type="outstation",
        base_package_rate=None,
        min_km_per_day=300.0,
        extra_km_rate=15.0,
        extra_hour_rate=None,
        driver_bata_rate=300.0
    )

    # Rates are distinct and not hardcoded to one global value
    assert windsor_sedan.base_package_rate != welspun_sedan.base_package_rate
    assert windsor_sedan.extra_km_rate != welspun_sedan.extra_km_rate
    assert windsor_sedan.extra_hour_rate != welspun_sedan.extra_hour_rate
    assert crysta_local.base_package_rate > windsor_sedan.base_package_rate

    # Ensure CompanyCreate supports billing_rates
    company_data = CompanyCreate(
        name="Windsor Park",
        code="WP01",
        billing_rates=[windsor_sedan, crysta_local, outstation_sedan]
    )
    assert len(company_data.billing_rates) == 3
    assert company_data.billing_rates[0].vehicle_category == "sedan"
    assert company_data.billing_rates[1].vehicle_category == "crysta"
    assert company_data.billing_rates[2].service_type == "outstation"


def test_generator_respects_passenger_and_booked_by_separation():
    """
    Test that WordBillGenerator uses passenger_name for 'For :' and booked_by for 'Booked by :'.
    CRITICAL: driver_name must NEVER appear in 'Booked by :'.
    """
    generator = WordBillGenerator()
    duty_slip = {
        "duty_slip_no": "501",
        "date": "05-09-2026",
        "driver_name": "Driver Mahesh",
        "passenger_name": "Guest Srikanth",
        "booked_by": "Versha Byte Solutions",
        "vehicle_number": "TS09CC5555 Sedan",
        "total_km": "100",
        "total_hours": "8.0",
        "starting_time": "09:00 AM",
        "closing_time": "05:00 PM",
        "base_package": "2500",
        "extra_km_charge": "300",
        "extra_hour_charge": "0",
        "bata": "0",
        "toll": "50",
        "parking": "100",
        "total_amount": "2950"
    }

    docx_bytes = generator.generate_single_bill_docx_bytes(
        duty_slip_no=duty_slip["duty_slip_no"],
        fields_dict=duty_slip,
        bill_no="B-501",
        company_name="Versha Byte Pvt Ltd"
    )
    doc = docx.Document(io.BytesIO(docx_bytes))
    text = "\n".join([p.text for p in doc.paragraphs])

    # Passenger must be in For : line
    assert "For :Guest Srikanth" in text

    # Booked by must be the client/booker, NOT the driver!
    assert "Booked by : Versha Byte Solutions" in text
    assert "Booked by : Driver Mahesh" not in text

    # Verify parking is included in the billing items table
    table = doc.tables[0]
    table_text = "\n".join([cell.text for row in table.rows for cell in row.cells])
    assert "Parking" in table_text
    assert "100" in table_text


def test_generator_fallback_when_booked_by_absent():
    """
    When booked_by is absent, WordBillGenerator should fall back to company_name,
    and NEVER driver_name.
    """
    generator = WordBillGenerator()
    duty_slip = {
        "duty_slip_no": "502",
        "driver_name": "Driver Santosh",
        "passenger_name": "Guest Vinay",
        "booked_by": None,
        "vehicle_number": "TS09DD6666 Sedan",
        "base_package": "2500",
        "total_amount": "2500"
    }

    docx_bytes = generator.generate_single_bill_docx_bytes(
        duty_slip_no=duty_slip["duty_slip_no"],
        fields_dict=duty_slip,
        bill_no="B-502",
        company_name="Wex Technologies"
    )
    doc = docx.Document(io.BytesIO(docx_bytes))
    text = "\n".join([p.text for p in doc.paragraphs])

    assert "Booked by : Wex Technologies" in text
    assert "Driver Santosh" not in text
