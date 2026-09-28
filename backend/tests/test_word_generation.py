import io
import docx
import pytest
from starlette.testclient import TestClient
from app.main import app
from app.core.database import db_manager
from app.core.config import settings
from app.services.word.num_to_words import amount_to_words_inr


@pytest.fixture(scope="module")
def client():
    original_provider = settings.OCR_PROVIDER
    settings.OCR_PROVIDER = "mock"
    with TestClient(app) as c:
        db = db_manager.get_database()
        # Clean up test records
        test_slips = list(db["duty_slips"].find({"duty_slip_no": {"$in": ["60", "61", "62", "63", "64"]}}))
        test_slip_ids = [s["_id"] for s in test_slips]
        db["extractions"].delete_many({"duty_slip_id": {"$in": test_slip_ids}})
        db["duty_slips"].delete_many({"duty_slip_no": {"$in": ["60", "61", "62", "63", "64"]}})
        db["companies"].delete_many({"name": {"$regex": "^TEST_WORD_"}})

        # Create test company A
        res_comp_a = c.post("/api/v1/companies", json={"name": "TEST_WORD_Company_A"})
        assert res_comp_a.status_code == 201
        company_a_id = res_comp_a.json()["id"]

        # Create test company B
        res_comp_b = c.post("/api/v1/companies", json={"name": "TEST_WORD_Company_B"})
        assert res_comp_b.status_code == 201
        company_b_id = res_comp_b.json()["id"]

        files = {
            "front_file": ("front.jpg", io.BytesIO(b"FRONT_IMAGE_60"), "image/jpeg"),
            "back_file": ("back.jpg", io.BytesIO(b"BACK_IMAGE_60"), "image/jpeg"),
        }
        # Company A Duty Slip #60
        res_ds60 = c.post("/api/v1/duty-slips", data={"company_id": company_a_id, "duty_slip_no": "60"}, files=files)
        assert res_ds60.status_code == 201
        ds60_id = res_ds60.json()["id"]

        # Company A Duty Slip #61 (Unverified)
        res_ds61 = c.post("/api/v1/duty-slips", data={"company_id": company_a_id, "duty_slip_no": "61"}, files=files)
        assert res_ds61.status_code == 201
        ds61_id = res_ds61.json()["id"]

        # Company A Duty Slip #62 (Second bill for same company)
        res_ds62 = c.post("/api/v1/duty-slips", data={"company_id": company_a_id, "duty_slip_no": "62"}, files=files)
        assert res_ds62.status_code == 201
        ds62_id = res_ds62.json()["id"]

        # Company B Duty Slip #63 (Different company)
        res_ds63 = c.post("/api/v1/duty-slips", data={"company_id": company_b_id, "duty_slip_no": "63"}, files=files)
        assert res_ds63.status_code == 201
        ds63_id = res_ds63.json()["id"]

        yield c, company_a_id, company_b_id, ds60_id, ds61_id, ds62_id, ds63_id

        # Teardown
        test_slips = list(db["duty_slips"].find({"duty_slip_no": {"$in": ["60", "61", "62", "63", "64"]}}))
        test_slip_ids = [s["_id"] for s in test_slips]
        db["extractions"].delete_many({"duty_slip_id": {"$in": test_slip_ids}})
        db["duty_slips"].delete_many({"duty_slip_no": {"$in": ["60", "61", "62", "63", "64"]}})
        db["companies"].delete_many({"name": {"$regex": "^TEST_WORD_"}})
        settings.OCR_PROVIDER = original_provider


def test_amount_in_words_inr():
    assert amount_to_words_inr(4020) == "Four Thousand Twenty Rupees Only"
    assert amount_to_words_inr(2500) == "Two Thousand Five Hundred Rupees Only"
    assert amount_to_words_inr(1005) == "One Thousand Five Rupees Only"
    assert amount_to_words_inr(0) == "Zero Rupees Only"
    assert amount_to_words_inr(125000) == "One Lakh Twenty Five Thousand Rupees Only"
    assert amount_to_words_inr(4020.50) == "Four Thousand Twenty Rupees and Fifty Paise Only"
    assert amount_to_words_inr(None) == ""


def test_unverified_extraction_blocks_word_generation(client):
    c, company_a_id, _, _, ds61_id, _, _ = client
    # 1. Run extraction (leaves status as extracted / ready_for_review)
    ext_res = c.post(f"/api/v1/duty-slips/{ds61_id}/extract")
    assert ext_res.status_code == 200

    # 2. Attempt Word bill generation on unverified slip -> must return 422
    res = c.post(f"/api/v1/duty-slips/{ds61_id}/generate-word")
    assert res.status_code == 422
    assert "must be 'verified'" in res.json()["detail"]


def test_missing_duty_slip_word_generation_returns_404(client):
    c, _, _, _, _, _, _ = client
    fake_id = "6a9ef59baf5bb217e24aa999"
    res = c.post(f"/api/v1/duty-slips/{fake_id}/generate-word")
    assert res.status_code == 404


def test_first_bill_creates_company_master_docx(client):
    c, company_a_id, _, ds60_id, _, _, _ = client

    # 1. Extract & Verify Duty Slip #60
    c.post(f"/api/v1/duty-slips/{ds60_id}/extract")
    c.put(
        f"/api/v1/duty-slips/{ds60_id}/extraction",
        json={
            "fields": {
                "vehicle_number": "TS09GC6246 A/c Sedan",
                "customer_name": "Uday Kumar",
                "extra_km": "67",
                "extra_hours": "1.5",
                "base_package": "8/80 = 2500",
                "extra_km_charge": "67 x 15 = 1005",
                "extra_hour_charge": "1.5 x 150 = 225",
                "bata": "250",
                "toll": "40",
                "total_amount": "4020"
            },
            "reviewer": "Reviewer_A"
        }
    )
    ver_res = c.post(f"/api/v1/duty-slips/{ds60_id}/verify")
    assert ver_res.status_code == 200

    # 2. Generate First Bill for Company A
    gen_res = c.post(f"/api/v1/duty-slips/{ds60_id}/generate-word", json={"bill_no": "Bill 01"})
    assert gen_res.status_code == 200
    data = gen_res.json()

    assert data["company_id"] == company_a_id
    assert data["bill_no"] == "Bill 01"
    assert data["filename"] == "TEST_WORD_Company_A.docx"
    assert data["total_bills_in_doc"] == 1
    assert data["total_amount"] == 4020.0
    assert data["amount_in_words"] == "Four Thousand Twenty Rupees Only"

    # 3. Download Company Master DOCX
    dl_res = c.get(f"/api/v1/companies/{company_a_id}/document/word")
    assert dl_res.status_code == 200
    assert dl_res.headers["content-type"] == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    assert 'attachment; filename="TEST_WORD_Company_A.docx"' in dl_res.headers["content-disposition"]

    docx_bytes = dl_res.content
    doc = docx.Document(io.BytesIO(docx_bytes))

    # A4 dimensions verification
    section = doc.sections[0]
    assert abs(section.page_width.mm - 210.0) < 1.0
    assert abs(section.page_height.mm - 297.0) < 1.0

    # Contains Bill 01 and reference elements
    all_text = "\n".join([p.text for p in doc.paragraphs])
    assert "Bill No.01" in all_text or "Bill 01" in all_text
    assert "SRI TULJA BHAVANI TRAVELS" in all_text
    assert "RENT-A-CAR" in all_text
    assert "Uday Kumar" in all_text
    assert "Rupees (in words): Four Thousand Twenty Rupees Only" in all_text

    # Verify table structure & cell contents
    assert len(doc.tables) == 1
    table = doc.tables[0]
    assert len(table.rows) == 3
    assert len(table.columns) == 9
    row1 = table.rows[1]
    assert "TS09GC6246 A/c Sedan" in row1.cells[2].text
    assert "67" in row1.cells[5].text
    assert "1.5" in row1.cells[6].text
    assert "8/80" in row1.cells[7].text
    assert "67 x 15" in row1.cells[7].text
    assert "1.5 x 150" in row1.cells[7].text
    assert "Bata" in row1.cells[7].text
    assert "Toll" in row1.cells[7].text
    assert "2500.00" in row1.cells[8].text
    assert "1005.00" in row1.cells[8].text
    assert "225.00" in row1.cells[8].text
    assert "250.00" in row1.cells[8].text
    assert "40.00" in row1.cells[8].text

    # Total row
    row2 = table.rows[2]
    assert "Total" in row2.cells[7].text
    assert "4020.00" in row2.cells[8].text


def test_second_bill_for_same_company_appends_to_master_docx(client):
    c, company_a_id, _, _, _, ds62_id, _ = client

    # 1. Extract & Verify Duty Slip #62 (belonging to Company A)
    c.post(f"/api/v1/duty-slips/{ds62_id}/extract")
    c.put(
        f"/api/v1/duty-slips/{ds62_id}/extraction",
        json={
            "fields": {
                "vehicle_number": "TS09GC6246 A/c Sedan",
                "customer_name": "Uday Kumar",
                "extra_km": "67",
                "extra_hours": "1.5",
                "base_package": "8/80 = 2500",
                "extra_km_charge": "67 x 15 = 1005",
                "extra_hour_charge": "1.5 x 150 = 225",
                "bata": "250",
                "toll": "40",
                "total_amount": "4020"
            },
            "reviewer": "Reviewer_A"
        }
    )
    c.post(f"/api/v1/duty-slips/{ds62_id}/verify")

    # 2. Generate Second Bill for Company A -> must append Bill 02 to existing document
    gen_res = c.post(f"/api/v1/duty-slips/{ds62_id}/generate-word", json={"bill_no": "Bill 02"})
    assert gen_res.status_code == 200
    data = gen_res.json()
    assert data["bill_no"] == "Bill 02"
    assert data["total_bills_in_doc"] == 2

    # 3. Download Company A Master DOCX and check for 2 bill tables & pages
    dl_res = c.get(f"/api/v1/companies/{company_a_id}/document/word")
    assert dl_res.status_code == 200

    doc = docx.Document(io.BytesIO(dl_res.content))
    all_text = "\n".join([p.text for p in doc.paragraphs])
    for table in doc.tables:
        for row in table.rows:
            all_text += "\n" + " | ".join([cell.text for cell in row.cells])

    # Must contain both Bill 01 and Bill 02 in the SAME document
    assert "Bill No.01" in all_text or "Bill 01" in all_text
    assert "Bill No.02" in all_text or "Bill 02" in all_text
    # Should have 2 billing tables for the 2 bills
    assert len(doc.tables) >= 2


def test_different_company_creates_isolated_master_docx(client):
    c, company_a_id, company_b_id, _, _, _, ds63_id = client

    # 1. Extract & Verify Duty Slip #63 (belonging to Company B)
    c.post(f"/api/v1/duty-slips/{ds63_id}/extract")
    c.put(
        f"/api/v1/duty-slips/{ds63_id}/extraction",
        json={
            "fields": {
                "duty_slip_no": "63",
                "vehicle_number": "AP09XY9999 Innova",
                "customer_name": "Company B Client",
                "starting_km": "1000",
                "closing_km": "1147",
                "total_km": "147",
                "base_package": "8/80 = 2500",
                "extra_km_charge": "67 x 15 = 1005",
                "extra_hour_charge": "1.5 x 150 = 225",
                "bata": "250",
                "toll": "40",
                "total_amount": "4020",
                "starting_time": "08:00 AM",
                "closing_time": "05:30 PM",
                "total_hours": "9.5",
                "date": "03-09-2026"
            },
            "reviewer": "Reviewer_B"
        }
    )
    ver_res = c.post(f"/api/v1/duty-slips/{ds63_id}/verify")
    assert ver_res.status_code == 200

    # 2. Generate Bill for Company B
    gen_res = c.post(f"/api/v1/duty-slips/{ds63_id}/generate-word", json={"bill_no": "Bill 01"})
    assert gen_res.status_code == 200
    data = gen_res.json()
    assert data["company_id"] == company_b_id
    assert data["filename"] == "TEST_WORD_Company_B.docx"
    assert data["total_bills_in_doc"] == 1

    # 3. Verify Company B document is isolated from Company A
    dl_b = c.get(f"/api/v1/companies/{company_b_id}/document/word")
    assert dl_b.status_code == 200
    doc_b = docx.Document(io.BytesIO(dl_b.content))
    text_b = "\n".join([p.text for p in doc_b.paragraphs])

    assert "Company B Client" in text_b
    assert "Uday Kumar" not in text_b  # Company A client must NOT be present in Company B document!


def test_company_master_info_endpoint(client):
    c, company_a_id, _, _, _, _, _ = client
    res = c.get(f"/api/v1/companies/{company_a_id}/document/info")
    assert res.status_code == 200
    info = res.json()
    assert info["has_master_doc"] is True
    assert info["total_bills"] == 2
    assert "TEST_WORD_Company_A.docx" in info["filename"]
    assert f"/api/v1/companies/{company_a_id}/document/word" == info["docx_url"]
    assert f"/api/v1/companies/{company_a_id}/document/pdf" == info["pdf_url"]


def test_company_master_pdf_endpoint(client):
    c, company_a_id, _, _, _, _, _ = client
    res = c.get(f"/api/v1/companies/{company_a_id}/document/pdf")
    # Should return PDF stream if COM is available
    if res.status_code == 200:
        assert res.headers["content-type"] == "application/pdf"
        assert len(res.content) > 0


def test_data_integrity_scans_and_audit_history_preserved(client):
    c, _, _, ds60_id, _, _, _ = client
    # 1. Scans are still present
    front_res = c.get(f"/api/v1/duty-slips/{ds60_id}/front")
    assert front_res.status_code == 200

    back_res = c.get(f"/api/v1/duty-slips/{ds60_id}/back")
    assert back_res.status_code == 200

    # 2. Extraction history and original OCR values are still intact
    ext_res = c.get(f"/api/v1/duty-slips/{ds60_id}/extraction")
    assert ext_res.status_code == 200
    ext_data = ext_res.json()
    assert ext_data["status"] == "verified"
    assert len(ext_data.get("audit_log", [])) >= 1
