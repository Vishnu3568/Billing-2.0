import requests
import os
import docx
import win32com.client
import fitz # PyMuPDF
import time

API_BASE = "http://127.0.0.1:8000/api/v1"
SLIP_ID = "6a9daac0d12cad7d3f26f9a7"
ARTIFACT_DIR = r"C:\Users\uvish\.gemini\antigravity-ide\brain\f0390443-2b92-48de-994f-b1770d6b57b6"

def main():
    print("=== LIVE END-TO-END STEP 7 REAL APP VERIFICATION ===")
    
    # 1. Check Duty Slip Status
    ds_res = requests.get(f"{API_BASE}/duty-slips/{SLIP_ID}")
    print(f"1. GET Duty Slip: HTTP {ds_res.status_code}")
    assert ds_res.status_code == 200, ds_res.text
    ds_data = ds_res.json()
    print(f"   Duty Slip No: {ds_data.get('duty_slip_no')}, Status: {ds_data.get('status')}")
    
    # 2. Run OCR Extraction
    print("2. Triggering OCR Extraction...")
    ext_run_res = requests.post(f"{API_BASE}/duty-slips/{SLIP_ID}/extract")
    print(f"   POST /extract: HTTP {ext_run_res.status_code}")
    assert ext_run_res.status_code == 200, ext_run_res.text
    ext_data = ext_run_res.json()
    print(f"   Extracted vehicle: {ext_data.get('extracted_fields', {}).get('vehicle_no')}")
    
    # 3. Apply Human Review Correction
    print("3. Applying Human Review Correction...")
    edit_payload = {
        "fields": {
            "vehicle_no": "TS09GC6243 A/C Sedan",
            "driver_name": "Ramu",
            "reporting_place": "Begumpet",
            "tour_location": "Local",
            "base_package": "8/80 = 2500",
            "extra_km_charge": "67 x 15 = 1005",
            "extra_hour_charge": "1.5 x 150 = 225",
            "bata": "250",
            "toll": "40",
            "total_amount": "4020",
            "starting_time": "08:00 AM",
            "closing_time": "05:30 PM",
            "meter_start": "1000",
            "meter_return": "1147",
            "total_km": "147",
            "total_hours": "9.5",
            "extra_km": "67",
            "extra_hours": "1.5",
            "date": "02-09-2026",
            "customer_name": "Uday Kumar"
        }
    }
    rev_res = requests.put(f"{API_BASE}/duty-slips/{SLIP_ID}/extraction", json=edit_payload)
    print(f"   PUT /extraction: HTTP {rev_res.status_code}")
    assert rev_res.status_code == 200, rev_res.text
    
    # 4. Verify Extraction
    print("4. Verifying Extraction...")
    ver_res = requests.post(f"{API_BASE}/duty-slips/{SLIP_ID}/verify")
    print(f"   POST /verify: HTTP {ver_res.status_code}")
    assert ver_res.status_code == 200, ver_res.text
    
    # 5. Trigger Real Generate Word Bill API
    print("5. Triggering Generate Word Bill...")
    gen_res = requests.post(f"{API_BASE}/duty-slips/{SLIP_ID}/generate-word", json={"bill_no": "BILL-2026-0047"})
    print(f"   POST Generate Word: HTTP {gen_res.status_code}")
    assert gen_res.status_code == 200, gen_res.text
    gen_data = gen_res.json()
    print(f"   Response filename: {gen_data.get('filename')}")
    print(f"   File size bytes: {gen_data.get('file_size_bytes')}")
    
    # 6. Download DOCX
    doc_res = requests.get(f"{API_BASE}/duty-slips/{SLIP_ID}/word-bill")
    print(f"6. GET Word Bill Download: HTTP {doc_res.status_code}")
    assert doc_res.status_code == 200, doc_res.text
    
    docx_path = os.path.join(ARTIFACT_DIR, "live_app_duty_slip_47_bill.docx")
    with open(docx_path, "wb") as f:
        f.write(doc_res.content)
    print(f"   Saved downloaded DOCX to: {docx_path} ({len(doc_res.content)} bytes)")
    
    # 7. Inspect DOCX Structure
    doc = docx.Document(docx_path)
    section = doc.sections[0]
    print(f"7. Page Dimensions:")
    print(f"   Width: {section.page_width.mm:.1f} mm (Expected ~210 mm A4)")
    print(f"   Height: {section.page_height.mm:.1f} mm (Expected ~297 mm A4)")
    assert abs(section.page_width.mm - 210) < 1.0, f"Width is not A4: {section.page_width.mm}"
    assert abs(section.page_height.mm - 297) < 1.0, f"Height is not A4: {section.page_height.mm}"
    
    print(f"   Paragraphs count: {len(doc.paragraphs)}")
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip():
            print(f"     P[{i}]: {p.text.strip()}")
            
    print(f"   Tables count: {len(doc.tables)}")
    for t_idx, table in enumerate(doc.tables):
        print(f"\n   --- Table {t_idx + 1} ({len(table.rows)} rows x {len(table.columns)} cols) ---")
        for r_idx, row in enumerate(table.rows):
            row_vals = [cell.text.strip().replace('\n', ' ') for cell in row.cells]
            print(f"     Row {r_idx}: {row_vals}")
            
    # 8. Render to PDF and PNG screenshot using MS Word COM
    print("\n8. Rendering via MS Word COM...")
    pdf_path = os.path.join(ARTIFACT_DIR, "live_app_duty_slip_47_bill.pdf")
    png_path = os.path.join(ARTIFACT_DIR, "actual_live_word_bill_47.png")
    
    word = win32com.client.Dispatch("Word.Application")
    word.Visible = False
    try:
        doc_com = word.Documents.Open(docx_path)
        doc_com.SaveAs(pdf_path, FileFormat=17)
        doc_com.Close(SaveChanges=0)
    finally:
        word.Quit()
        
    time.sleep(1)
    
    pdf_doc = fitz.open(pdf_path)
    page = pdf_doc[0]
    pix = page.get_pixmap(dpi=300)
    pix.save(png_path)
    pdf_doc.close()
    
    print(f"   Successfully generated rendered PNG: {png_path} ({os.path.getsize(png_path)} bytes)")
    print("=== LIVE VERIFICATION COMPLETE ===")

if __name__ == "__main__":
    main()
