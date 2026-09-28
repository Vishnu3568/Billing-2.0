import io
from typing import Dict, Any, Optional, List, Tuple
import docx
from docx.shared import Inches, Pt, Mm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

from app.services.word.num_to_words import amount_to_words_inr


FONT_PRIMARY = "Bookman Old Style"
FONT_TITLE = "Imprint MT Shadow"


def _set_cell_margins(cell, top=80, bottom=80, left=100, right=100):
    """Set inner margins (padding) for a table cell in dxa (1 pt = 20 dxa)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)


def _set_cell_border(cell, top="single", bottom="single", left="single", right="single", sz="4", color="333333"):
    """Set individual cell borders."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="{top}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:left w:val="{left}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="{bottom}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:right w:val="{right}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(tcBorders)


def _set_section_page_borders(section):
    """Set decorative outer page border on section matching reference document."""
    sectPr = section._sectPr
    # Remove existing pgBorders if present
    for existing in sectPr.xpath('./w:pgBorders'):
        sectPr.remove(existing)
    
    pg_borders = parse_xml(
        f'<w:pgBorders {nsdecls("w")} w:offsetFrom="page">'
        f'<w:top w:val="thinThickSmallGap" w:sz="24" w:space="27" w:color="auto"/>'
        f'<w:left w:val="thinThickSmallGap" w:sz="24" w:space="27" w:color="auto"/>'
        f'<w:bottom w:val="thickThinSmallGap" w:sz="24" w:space="27" w:color="auto"/>'
        f'<w:right w:val="thickThinSmallGap" w:sz="24" w:space="27" w:color="auto"/>'
        f'</w:pgBorders>'
    )
    sectPr.append(pg_borders)


def _set_row_height(row, height_dxa=2663, hRule="atLeast"):
    """Set explicit row height in dxa (1 pt = 20 dxa)."""
    trPr = row._tr.get_or_add_trPr()
    for existing in trPr.xpath('./w:trHeight'):
        trPr.remove(existing)
    tr_height = parse_xml(
        f'<w:trHeight {nsdecls("w")} w:val="{height_dxa}" w:hRule="{hRule}"/>'
    )
    trPr.append(tr_height)


class WordBillGenerator:
    """
    Dedicated generator for Company Master Word Documents (.docx)
    matching the exact reference layout, A4 page size, multi-page bill structure,
    and reference 9-column billing table.
    """

    @classmethod
    def create_or_append_bill_to_master(
        cls,
        existing_docx_bytes: Optional[bytes],
        duty_slip_no: str,
        fields_dict: Dict[str, Any],
        company_name: Optional[str] = None,
        bill_no: Optional[str] = None,
        total_hours: Optional[float] = None,
    ) -> bytes:
        """
        Creates a new Company Master DOCX or opens an existing one,
        appends a new Bill page (with page break), and returns the updated DOCX bytes.
        """
        if existing_docx_bytes:
            doc = docx.Document(io.BytesIO(existing_docx_bytes))
            # Ensure page border on all sections
            for section in doc.sections:
                _set_section_page_borders(section)
            # If document already has content, add a page break before appending new bill
            if len(doc.paragraphs) > 0 or len(doc.tables) > 0:
                doc.add_page_break()
        else:
            doc = docx.Document()
            # Set up page dimensions & borders for the master document
            for section in doc.sections:
                section.page_width = Mm(210)    # ISO 216 A4 (210mm x 297mm)
                section.page_height = Mm(297)
                section.top_margin = Inches(0.5)
                section.bottom_margin = Inches(0.5)
                section.left_margin = Inches(0.65)
                section.right_margin = Inches(0.65)
                _set_section_page_borders(section)

        cls._render_single_bill_page(
            doc=doc,
            duty_slip_no=duty_slip_no,
            fields_dict=fields_dict,
            company_name=company_name,
            bill_no=bill_no,
            total_hours=total_hours,
        )

        out_io = io.BytesIO()
        doc.save(out_io)
        return out_io.getvalue()

    @classmethod
    def generate_single_bill_docx_bytes(
        cls,
        duty_slip_no: str,
        fields_dict: Dict[str, Any],
        company_name: Optional[str] = None,
        bill_no: Optional[str] = None,
        total_hours: Optional[float] = None,
    ) -> bytes:
        """Generates an individual standalone bill document."""
        return cls.create_or_append_bill_to_master(
            existing_docx_bytes=None,
            duty_slip_no=duty_slip_no,
            fields_dict=fields_dict,
            company_name=company_name,
            bill_no=bill_no,
            total_hours=total_hours,
        )

    @classmethod
    def _render_single_bill_page(
        cls,
        doc: docx.Document,
        duty_slip_no: str,
        fields_dict: Dict[str, Any],
        company_name: Optional[str] = None,
        bill_no: Optional[str] = None,
        total_hours: Optional[float] = None,
    ) -> None:
        """Renders one complete Bill page into the document."""
        # Helper to format field value
        def get_val(fname: str, default: str = "") -> str:
            f = fields_dict.get(fname)
            if f is None:
                return default
            if isinstance(f, dict):
                v = f.get("value")
                return str(v) if v is not None else default
            return str(f) if f is not None else default

        # Extract verified values
        date_val = get_val("date", "")
        driver_val = get_val("driver_name", "")
        vehicle_val = get_val("vehicle_number") or get_val("vehicle_no", "")
        party_val = get_val("customer_name") or get_val("party_name", "") or company_name or ""
        total_km_val = get_val("total_km", "")
        starting_time_val = get_val("starting_time", "")
        closing_time_val = get_val("closing_time", "")
        extra_km_val = get_val("extra_km", "")
        extra_hours_val = get_val("extra_hours", "")

        # Billing breakdown items
        base_pkg_raw = get_val("billing_base_package", "") or get_val("base_package", "")
        extra_km_raw = get_val("billing_extra_km", "") or get_val("extra_km_charge", "")
        extra_hrs_raw = get_val("billing_extra_hours", "") or get_val("extra_hour_charge", "")
        bata_raw = get_val("bata", "")
        toll_raw = get_val("toll", "")
        total_amt_raw = get_val("total_amount", "")

        # Determine total hours display
        total_hrs_str = ""
        if total_hours is not None:
            total_hrs_str = f"{total_hours:g}"
        elif starting_time_val and closing_time_val:
            total_hrs_str = f"{starting_time_val} to {closing_time_val}"
        else:
            total_hrs_str = get_val("total_hours", "")

        # Display Bill No.
        display_bill_no = bill_no if bill_no else (duty_slip_no if duty_slip_no else "01")
        if not display_bill_no.lower().startswith("bill"):
            # Ensure 2 digits if numeric
            if display_bill_no.isdigit() and len(display_bill_no) == 1:
                display_bill_no = f"0{display_bill_no}"
            display_bill_no_text = f"Bill No.{display_bill_no}"
        else:
            display_bill_no_text = display_bill_no

        # ----------------------------------------------------
        # 1. HEADER SECTION (Faithful Reference Match)
        # ----------------------------------------------------
        # Phone numbers: Imprint MT Shadow, 11pt, #0070C0, Center aligned
        p_phone = doc.add_paragraph()
        p_phone.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_phone.paragraph_format.space_before = Pt(0)
        p_phone.paragraph_format.space_after = Pt(2)
        r_phone = p_phone.add_run("Mobile No: 94405 22 814, 99892 08711, 9000 240 410")
        r_phone.font.name = FONT_TITLE
        r_phone.font.size = Pt(11)
        r_phone.font.bold = False
        r_phone.font.color.rgb = RGBColor(0x00, 0x70, 0xC0)  # Reference Blue #0070C0

        # Company Title: Imprint MT Shadow, 27pt, #002060 (Navy Blue), Center aligned
        p_name = doc.add_paragraph()
        p_name.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_name.paragraph_format.space_before = Pt(0)
        p_name.paragraph_format.space_after = Pt(1)
        r_name = p_name.add_run("SRI TULJA BHAVANI TRAVELS")
        r_name.font.name = FONT_TITLE
        r_name.font.size = Pt(27)
        r_name.font.bold = True
        r_name.font.color.rgb = RGBColor(0x00, 0x20, 0x60)  # Navy Blue #002060

        # Subtitle: Imprint MT Shadow, 20pt, #FF0000 (Red), Center aligned
        p_sub = doc.add_paragraph()
        p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_sub.paragraph_format.space_before = Pt(0)
        p_sub.paragraph_format.space_after = Pt(2)
        r_sub = p_sub.add_run("RENT-A-CAR")
        r_sub.font.name = FONT_TITLE
        r_sub.font.size = Pt(20)
        r_sub.font.bold = True
        r_sub.font.color.rgb = RGBColor(0xFF, 0x00, 0x00)  # Pure Red #FF0000

        # Address & Email: Imprint MT Shadow, 11pt, #002060, Center aligned
        p_addr = doc.add_paragraph()
        p_addr.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_addr.paragraph_format.space_before = Pt(0)
        p_addr.paragraph_format.space_after = Pt(4)
        r_addr = p_addr.add_run("1-11-113/3,P2 Sai Shikara Apartments, Shayamlal Building Begumpet, Hyderabad - 500016, srituljabhavanitravels.rentacar@gmail.com")
        r_addr.font.name = FONT_TITLE
        r_addr.font.size = Pt(11)
        r_addr.font.color.rgb = RGBColor(0x00, 0x20, 0x60)

        # Horizontal separator line beneath header
        pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="12" w:space="4" w:color="002060"/></w:pBdr>')
        p_addr._p.get_or_add_pPr().append(pBdr)

        # ----------------------------------------------------
        # 2. BILL METADATA & RECIPIENT BLOCK
        # ----------------------------------------------------
        # Bill No & Date paragraph with tab alignment (Bookman Old Style 11pt)
        p_bill_date = doc.add_paragraph()
        p_bill_date.paragraph_format.space_before = Pt(8)
        p_bill_date.paragraph_format.space_after = Pt(4)
        r_b = p_bill_date.add_run(f"{display_bill_no_text}")
        r_b.font.name = FONT_PRIMARY
        r_b.font.size = Pt(11)
        r_b.font.bold = False

        r_sp = p_bill_date.add_run("\t\t\t\t\t\t             ")
        r_d = p_bill_date.add_run(f"Date: {date_val}")
        r_d.font.name = FONT_PRIMARY
        r_d.font.size = Pt(11)
        r_d.font.bold = False

        p_to = doc.add_paragraph()
        p_to.paragraph_format.space_before = Pt(4)
        p_to.paragraph_format.space_after = Pt(2)
        r_to = p_to.add_run("To, ")
        r_to.font.name = FONT_PRIMARY
        r_to.font.size = Pt(11)
        r_to.font.bold = False

        p_cust = doc.add_paragraph()
        p_cust.paragraph_format.space_before = Pt(0)
        p_cust.paragraph_format.space_after = Pt(10)
        r_cust = p_cust.add_run(f"{party_val}")
        r_cust.font.name = FONT_PRIMARY
        r_cust.font.size = Pt(11)
        r_cust.font.bold = False

        # ----------------------------------------------------
        # 3. 9-COLUMN BILLING TABLE (Exact Reference Structure)
        # ----------------------------------------------------
        # Row 0: Headers
        # Row 1: Multi-line Data items
        # Row 2: Total
        table = doc.add_table(rows=3, cols=9)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False

        headers = [
            "Duty Slip No.", "Date", "Vehicle No.", "Total Kms",
            "Total Hrs.", "Extra Kms.", "Extra Hrs.", "Amt", "Total Amount"
        ]
        col_widths = [
            Inches(0.75), Inches(0.85), Inches(1.3), Inches(0.7),
            Inches(0.7), Inches(0.75), Inches(0.9), Inches(0.75), Inches(0.95)
        ]

        # Row 0: Header Row
        hdr_row = table.rows[0]
        _set_row_height(hdr_row, height_dxa=770, hRule="atLeast")
        for idx, text in enumerate(headers):
            cell = hdr_row.cells[idx]
            cell.width = col_widths[idx]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            _set_cell_margins(cell, top=80, bottom=80, left=60, right=60)
            _set_cell_border(cell, top="single", bottom="single", left="single", right="single", sz="4", color="000000")
            
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(text)
            run.font.name = FONT_PRIMARY
            run.font.size = Pt(9.0)
            run.font.bold = True

        # Parse amounts and breakdown
        def parse_amt(val_str: Any) -> Optional[float]:
            if val_str is None or val_str == "":
                return None
            try:
                s = str(val_str).strip()
                if "=" in s:
                    s = s.split("=")[-1].strip()
                cleaned = "".join(c for c in s if c.isdigit() or c == ".")
                return float(cleaned) if cleaned else None
            except ValueError:
                return None

        def extract_expr(val: Any, default_label: str) -> str:
            if val is None or val == "":
                return default_label
            s = str(val).strip()
            if "=" in s:
                return s.split("=")[0].strip()
            return default_label

        # ----------------------------------------------------
        # BILLING LINE ITEMS (Amt column expressions & Total Amount values)
        # ----------------------------------------------------
        billing_items: List[Tuple[str, float]] = []

        # 1. Base Package (e.g. "8/80 = 2500" -> expr "8/80", amt 2500.00)
        base_amt = parse_amt(base_pkg_raw)
        if base_amt is not None:
            base_expr = extract_expr(base_pkg_raw, "8/80")
            billing_items.append((base_expr, base_amt))

        # 2. Extra Kms (e.g. "67 x 15 = 1005" -> expr "67 x 15", amt 1005.00)
        extra_km_amt = parse_amt(extra_km_raw)
        if extra_km_amt is not None and extra_km_amt > 0:
            extra_km_expr = extract_expr(extra_km_raw, f"{extra_km_val} x 15" if extra_km_val else "Extra Km")
            billing_items.append((extra_km_expr, extra_km_amt))

        # 3. Extra Hours (e.g. "1.5 x 150 = 225" -> expr "1.5 x 150", amt 225.00)
        extra_hrs_amt = parse_amt(extra_hrs_raw)
        if extra_hrs_amt is not None and extra_hrs_amt > 0:
            extra_hrs_expr = extract_expr(extra_hrs_raw, f"{extra_hours_val} x 150" if extra_hours_val else "Extra Hrs")
            billing_items.append((extra_hrs_expr, extra_hrs_amt))

        # 4. Bata (e.g. "Bata = 250" or "250" -> expr "Bata", amt 250.00)
        bata_amt = parse_amt(bata_raw)
        if bata_amt is not None and bata_amt > 0:
            bata_expr = extract_expr(bata_raw, "Bata") if ("bata" in str(bata_raw).lower() and "=" in str(bata_raw)) else "Bata"
            billing_items.append((bata_expr, bata_amt))

        # 5. Toll (e.g. "Toll = 40" or "40" -> expr "Toll", amt 40.00)
        toll_amt = parse_amt(toll_raw)
        if toll_amt is not None and toll_amt > 0:
            toll_expr = extract_expr(toll_raw, "Toll") if ("toll" in str(toll_raw).lower() and "=" in str(toll_raw)) else "Toll"
            billing_items.append((toll_expr, toll_amt))

        # 6. Total Amount
        total_amt = parse_amt(total_amt_raw)
        if total_amt is None and billing_items:
            total_amt = sum(item[1] for item in billing_items)

        # Build column lines
        amt_col_text = "\n".join(item[0] for item in billing_items)
        tot_amt_col_text = "\n".join(f"{item[1]:.2f}" for item in billing_items)

        # Row 1: Data Row (with open minimum row height matching reference)
        data_row = table.rows[1]
        _set_row_height(data_row, height_dxa=2663, hRule="atLeast")
        cell_contents = [
            duty_slip_no,
            date_val,
            vehicle_val,
            total_km_val,
            total_hrs_str,
            extra_km_val,
            str(extra_hours_val) if extra_hours_val else "",
            amt_col_text,
            tot_amt_col_text
        ]

        for idx, content in enumerate(cell_contents):
            cell = data_row.cells[idx]
            cell.width = col_widths[idx]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP
            _set_cell_margins(cell, top=100, bottom=100, left=60, right=60)
            _set_cell_border(cell, top="single", bottom="single", left="single", right="single", sz="4", color="000000")
            
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT if idx == 8 else (WD_ALIGN_PARAGRAPH.CENTER if idx in [0, 1, 3, 4, 5, 6] else WD_ALIGN_PARAGRAPH.LEFT)
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            
            # Format lines with consistent line spacing
            lines = str(content).split("\n")
            for l_idx, line in enumerate(lines):
                if l_idx > 0:
                    p = cell.add_paragraph()
                    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT if idx == 8 else (WD_ALIGN_PARAGRAPH.CENTER if idx in [0, 1, 3, 4, 5, 6] else WD_ALIGN_PARAGRAPH.LEFT)
                    p.paragraph_format.space_before = Pt(2)
                    p.paragraph_format.space_after = Pt(0)
                run = p.add_run(line)
                run.font.name = FONT_PRIMARY
                run.font.size = Pt(9.5)

        # Row 2: Total Row
        tot_row = table.rows[2]
        _set_row_height(tot_row, height_dxa=620, hRule="atLeast")
        for idx in range(9):
            cell = tot_row.cells[idx]
            cell.width = col_widths[idx]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            _set_cell_margins(cell, top=80, bottom=80, left=60, right=60)
            _set_cell_border(cell, top="single", bottom="single", left="single", right="single", sz="4", color="000000")
            
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            
            if idx == 8:
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                tot_display = f"{total_amt:.2f}" if total_amt is not None else total_amt_raw
                run = p.add_run(tot_display)
                run.font.name = FONT_PRIMARY
                run.font.size = Pt(10)
                run.font.bold = True
            elif idx == 7:
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                run = p.add_run("Total")
                run.font.name = FONT_PRIMARY
                run.font.size = Pt(10)
                run.font.bold = True

        # ----------------------------------------------------
        # 4. FOOTER & SIGN-OFF SECTION (Matching Reference Typography)
        # ----------------------------------------------------
        # Amount in words: Imprint MT Shadow, 14pt
        words_text = ""
        if total_amt is not None:
            words_text = amount_to_words_inr(total_amt)
        else:
            words_text = f"{total_amt_raw} only"

        p_words = doc.add_paragraph()
        p_words.paragraph_format.space_before = Pt(14)
        p_words.paragraph_format.space_after = Pt(10)
        r_w = p_words.add_run(f"Rupees (in words): {words_text}")
        r_w.font.name = FONT_TITLE
        r_w.font.size = Pt(14)
        r_w.font.bold = False

        # Sign-off Header: Imprint MT Shadow, 14pt, Right-aligned
        p_sign_hdr = doc.add_paragraph()
        p_sign_hdr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p_sign_hdr.paragraph_format.space_before = Pt(6)
        p_sign_hdr.paragraph_format.space_after = Pt(4)
        r_sign_hdr = p_sign_hdr.add_run("For Sri Tulja Bhavani Travels")
        r_sign_hdr.font.name = FONT_TITLE
        r_sign_hdr.font.size = Pt(14)
        r_sign_hdr.font.bold = False

        # For recipient line: Imprint MT Shadow, 11pt
        p_for = doc.add_paragraph()
        p_for.paragraph_format.space_before = Pt(4)
        p_for.paragraph_format.space_after = Pt(16)
        r_for = p_for.add_run(f"For :{party_val}")
        r_for.font.name = FONT_TITLE
        r_for.font.size = Pt(11)
        r_for.font.bold = False

        # Booked by and Manager sign-off paragraph with tab alignment
        p_sign = doc.add_paragraph()
        p_sign.paragraph_format.space_before = Pt(8)
        p_sign.paragraph_format.space_after = Pt(0)
        booked_text = f"Booked by : {driver_val}" if driver_val else "Booked by :"
        r_bk = p_sign.add_run(booked_text)
        r_bk.font.name = FONT_TITLE
        r_bk.font.size = Pt(11)

        r_sp = p_sign.add_run("\t\t\t\t\t\t\t\t")
        r_mgr = p_sign.add_run("Manager")
        r_mgr.font.name = FONT_TITLE
        r_mgr.font.size = Pt(11)
        r_mgr.font.bold = True
