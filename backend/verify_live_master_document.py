import requests
import os
import docx
import win32com.client
import fitz # PyMuPDF
import time

API_BASE = "http://127.0.0.1:8000/api/v1"
SLIP_ID_47 = "6a9daac0d12cad7d3f26f9a7"
COMPANY_ID = "6a9daac0d12cad7d3f26f9a6"
ARTIFACT_DIR = r"C:\Users\uvish\.gemini\antigravity-ide\brain\f0390443-2b92-48de-994f-b1770d6b57b6"

def main():
    print("=== LIVE END-TO-END COMPANY MASTER WORD DOCUMENT VERIFICATION ===")
    
    # 1. Check Duty Slip #47
    ds_res = requests.get(f"{API_BASE}/duty-slips/{SLIP_ID_47}")
    print(f"1. GET Duty Slip #47: HTTP {ds_res.status_code}")
    assert ds_res.status_code == 200, ds_res.text
    ds_data = ds_res.json()
    print(f"   Company: {ds_data.get('company_name')} ({ds_data.get('company_id')})")
    print(f"   Duty Slip No: {ds_data.get('duty_slip_no')}, Status: {ds_data.get('status')}")

    # Reset company master document storage & db tracking for fresh clean test
    import pymongo
    from bson import ObjectId
    import shutil
    client = pymongo.MongoClient("mongodb://localhost:27017")
    db = client["billing_db"]
    db["companies"].update_one(
        {"_id": ObjectId(COMPANY_ID)},
        {"$unset": {"word_bills": "", "master_doc_path": ""}}
    )
    db["duty_slips"].update_many(
        {"company_id": ObjectId(COMPANY_ID)},
        {"$unset": {"word_bill": "", "has_word_bill": ""}}
    )
    company_master_dir = os.path.join(os.path.dirname(__file__), "storage", "companies", COMPANY_ID, "master")
    if os.path.exists(company_master_dir):
        shutil.rmtree(company_master_dir, ignore_errors=True)



    # 2. Extract, Review, and Verify Duty Slip #47
    print("2. Ensuring extraction is verified for Duty Slip #47...")
    requests.post(f"{API_BASE}/duty-slips/{SLIP_ID_47}/extract")
    requests.put(f"{API_BASE}/duty-slips/{SLIP_ID_47}/extraction", json={
        "fields": {
            "duty_slip_no": "47",
            "vehicle_number": "TS09GC6246 A/c Sedan",
            "driver_name": "Ramu",
            "customer_name": "Uday Kumar",
            "reporting_place": "Begumpet",
            "tour_location": "Local",
            "starting_km": "1000",
            "closing_km": "1147",
            "total_km": "147",
            "starting_time": "08:00 AM",
            "closing_time": "05:30 PM",
            "total_hours": "9.5",
            "extra_km": "67",
            "extra_hours": "1.5",
            "base_package": "8/80 = 2500",
            "extra_km_charge": "67 x 15 = 1005",
            "extra_hour_charge": "1.5 x 150 = 225",
            "bata": "250",
            "toll": "40",
            "total_amount": "4020",
            "date": "02-09-2026"
        },
        "reviewer": "Reviewer_Step7_Live"
    })
    ver_res = requests.post(f"{API_BASE}/duty-slips/{SLIP_ID_47}/verify")
    print(f"   POST /verify: HTTP {ver_res.status_code}")
    assert ver_res.status_code == 200, ver_res.text

    # 3. Generate Bill 01 for Duty Slip #47 -> Creates Company Master Document
    print("3. Generating Bill 01 for Company A Master Document...")
    gen_res = requests.post(f"{API_BASE}/duty-slips/{SLIP_ID_47}/generate-word", json={"bill_no": "Bill 01"})
    print(f"   POST /generate-word: HTTP {gen_res.status_code}")
    assert gen_res.status_code == 200, gen_res.text
    gen_data = gen_res.json()
    print(f"   Master Filename: {gen_data.get('filename')}")
    print(f"   Total Bills in Document: {gen_data.get('total_bills_in_doc')}")
    print(f"   Master DOCX URL: {gen_data.get('company_master_doc_url')}")
    print(f"   PDF Preview URL: {gen_data.get('pdf_preview_url')}")

    # 4. Create and Verify Duty Slip #48 for the SAME company to test multi-page appending
    print("\n4. Creating & verifying second Duty Slip #48 for same company...", flush=True)
    all_slips = requests.get(f"{API_BASE}/duty-slips?company_id={COMPANY_ID}").json()
    slip_48 = next((s for s in all_slips if s["duty_slip_no"] == "48"), None)
    
    if not slip_48:
        real_front_bytes = requests.get(f"{API_BASE}/duty-slips/{SLIP_ID_47}/front").content
        real_back_bytes = requests.get(f"{API_BASE}/duty-slips/{SLIP_ID_47}/back").content
        create_res = requests.post(
            f"{API_BASE}/duty-slips",
            data={"company_id": COMPANY_ID, "duty_slip_no": "48"},
            files={
                "front_file": ("front_48.jpeg", real_front_bytes, "image/jpeg"),
                "back_file": ("back_48.jpeg", real_back_bytes, "image/jpeg")
            }
        )
        assert create_res.status_code == 201, create_res.text
        slip_48 = create_res.json()
        requests.post(f"{API_BASE}/duty-slips/{slip_48['id']}/extract")

    slip_48_id = slip_48["id"]
    print(f"   Duty Slip 48 ID: {slip_48_id}", flush=True)

    # Ensure extraction exists and save verified fields
    requests.put(f"{API_BASE}/duty-slips/{slip_48_id}/extraction", json={
        "fields": {
            "duty_slip_no": "48",
            "vehicle_number": "TS09GC6246 A/c Sedan",
            "driver_name": "Ramu",
            "customer_name": "Uday Kumar",
            "starting_km": "1147",
            "closing_km": "1294",
            "total_km": "147",
            "starting_time": "08:00 AM",
            "closing_time": "05:30 PM",
            "total_hours": "9.5",
            "extra_km": "67",
            "extra_hours": "1.5",
            "base_package": "8/80 = 2500",
            "extra_km_charge": "67 x 15 = 1005",
            "extra_hour_charge": "1.5 x 150 = 225",
            "bata": "250",
            "toll": "40",
            "total_amount": "4020",
            "date": "03-09-2026"
        },
        "reviewer": "Reviewer_Step7_Live"
    })
    v48_res = requests.post(f"{API_BASE}/duty-slips/{slip_48_id}/verify")
    print(f"   Duty Slip 48 verified: HTTP {v48_res.status_code}", flush=True)

    # 5. Generate Bill 02 -> Appends to existing Company Master Document!
    print("5. Generating Bill 02 for same company (Appending to Master Document)...", flush=True)
    gen2_res = requests.post(f"{API_BASE}/duty-slips/{slip_48_id}/generate-word", json={"bill_no": "Bill 02"})
    print(f"   POST /generate-word (Bill 02): HTTP {gen2_res.status_code}", flush=True)
    assert gen2_res.status_code == 200, gen2_res.text
    gen2_data = gen2_res.json()
    print(f"   Master Filename: {gen2_data.get('filename')}", flush=True)
    print(f"   Total Bills in Document: {gen2_data.get('total_bills_in_doc')} (Expected 2)", flush=True)
    assert gen2_data.get('total_bills_in_doc') == 2


    # 6. Download Company Master Document (.docx)
    print("\n6. Downloading authoritative Company Master DOCX...")
    dl_res = requests.get(f"{API_BASE}/companies/{COMPANY_ID}/document/word")
    print(f"   GET /companies/{COMPANY_ID}/document/word: HTTP {dl_res.status_code}")
    assert dl_res.status_code == 200, dl_res.text

    master_docx_path = os.path.join(ARTIFACT_DIR, "live_company_master_document.docx")
    with open(master_docx_path, "wb") as f:
        f.write(dl_res.content)
    print(f"   Saved Master DOCX: {master_docx_path} ({len(dl_res.content)} bytes)")

    # 7. Inspect Master DOCX Structure
    doc = docx.Document(master_docx_path)
    print(f"7. Master Document Structure:")
    print(f"   Paragraphs: {len(doc.paragraphs)}")
    print(f"   Tables: {len(doc.tables)} (Expected 2 main billing tables for 2 bills)")
    
    all_text = "\n".join([p.text for p in doc.paragraphs])
    for t_idx, table in enumerate(doc.tables):
        print(f"\n   --- Billing Table {t_idx + 1} ({len(table.rows)} rows x {len(table.columns)} cols) ---")
        for r_idx, row in enumerate(table.rows):
            row_vals = [c.text.strip().replace('\n', ' | ') for c in row.cells]
            print(f"     Row {r_idx}: {row_vals}")

    # Check that both Bill 01 and Bill 02 are present
    assert "Bill No.01" in all_text or "Bill 01" in all_text, "Bill 01 missing!"
    assert "Bill No.02" in all_text or "Bill 02" in all_text, "Bill 02 missing!"
    assert len(doc.tables) == 2, f"Expected 2 tables, got {len(doc.tables)}"

    # 8. Render Company Master Document to Multi-page PDF and PNG Screenshots
    print("\n8. Rendering multi-page PDF via MS Word COM...")
    pdf_path = os.path.join(ARTIFACT_DIR, "live_company_master_document.pdf")
    
    word = win32com.client.Dispatch("Word.Application")
    word.Visible = False
    try:
        doc_com = word.Documents.Open(master_docx_path)
        doc_com.SaveAs(pdf_path, FileFormat=17) # 17 = wdFormatPDF
        doc_com.Close(SaveChanges=0)
    finally:
        word.Quit()

    time.sleep(1)

    # Convert each PDF page to PNG screenshot
    pdf_doc = fitz.open(pdf_path)
    print(f"   Rendered PDF Page Count: {len(pdf_doc)} pages")
    
    for p_num in range(len(pdf_doc)):
        page = pdf_doc[p_num]
        pix = page.get_pixmap(dpi=300)
        img_path = os.path.join(ARTIFACT_DIR, f"live_master_page_{p_num + 1}.png")
        pix.save(img_path)
        print(f"   Saved Page {p_num + 1} Screenshot: {img_path} ({os.path.getsize(img_path)} bytes)")
    
    pdf_doc.close()

    # 9. Test In-Browser PDF Stream Endpoint
    print("\n9. Testing GET /companies/{company_id}/document/pdf stream...")
    pdf_stream_res = requests.get(f"{API_BASE}/companies/{COMPANY_ID}/document/pdf")
    print(f"   GET /companies/{COMPANY_ID}/document/pdf: HTTP {pdf_stream_res.status_code}")
    assert pdf_stream_res.status_code == 200
    assert pdf_stream_res.headers["content-type"] == "application/pdf"
    print(f"   Streamed PDF size: {len(pdf_stream_res.content)} bytes")

    print("\n=== LIVE COMPANY MASTER VERIFICATION COMPLETE ===")

if __name__ == "__main__":
    main()
