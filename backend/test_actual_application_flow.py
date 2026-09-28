import io
import requests
import fitz  # PyMuPDF
import win32com.client
from pathlib import Path
from docx import Document

BASE_URL = "http://127.0.0.1:8000/api/v1"
DUTY_SLIP_ID = "6a9daac0d12cad7d3f26f9a7"


def run_actual_live_flow():
    print(f"=== TESTING ACTUAL LIVE APPLICATION FLOW FOR DUTY SLIP {DUTY_SLIP_ID} ===")

    # 1. GET Duty Slip
    print("\n[1] GET Duty Slip:")
    res_ds = requests.get(f"{BASE_URL}/duty-slips/{DUTY_SLIP_ID}")
    print(f"Status: {res_ds.status_code}")
    assert res_ds.status_code == 200, f"Failed to get duty slip: {res_ds.text}"
    ds_data = res_ds.json()
    print(f"Duty Slip No: {ds_data['duty_slip_no']}, Company: {ds_data['company_name']}")
    print(f"Front Scan Attached: {ds_data['front_scan']['original_filename']}")
    print(f"Back Scan Attached: {ds_data['back_scan']['original_filename']}")

    # 2. Check or Run OCR Extraction on real scans
    print("\n[2] Checking Extraction Result (from Live OCR):")
    res_ext = requests.get(f"{BASE_URL}/duty-slips/{DUTY_SLIP_ID}/extraction")
    if res_ext.status_code != 200:
        res_ext = requests.post(f"{BASE_URL}/duty-slips/{DUTY_SLIP_ID}/extract")
    assert res_ext.status_code == 200, f"Extraction failed: {res_ext.text}"
    ext_data = res_ext.json()
    print(f"Extraction Status: {ext_data['status']}")
    print(f"Overall Confidence: {ext_data['overall_confidence']}")
    print(f"Raw Extracted Vehicle: {ext_data['fields']['vehicle_number']['value']}")
    print(f"Raw Extracted Total Amount: {ext_data['fields']['total_amount']['value']}")
    print(f"Raw Extracted Line Items:")
    print(f"  Base Package: {ext_data['fields'].get('base_package', {}).get('value')}")
    print(f"  Extra KM Charge: {ext_data['fields'].get('extra_km_charge', {}).get('value')}")
    print(f"  Extra Hour Charge: {ext_data['fields'].get('extra_hour_charge', {}).get('value')}")
    print(f"  Bata: {ext_data['fields'].get('bata', {}).get('value')}")
    print(f"  Toll: {ext_data['fields'].get('toll', {}).get('value')}")

    # 3. Human Review & Editing: Ensure vehicle number is corrected to 'TS09GC6243 A/C Sedan'
    print("\n[3] PUT Human Review Corrections:")
    review_payload = {
        "fields": {
            "vehicle_number": "TS09GC6243 A/C Sedan",
            "customer_name": "Uday Kumar",
            "date": "02-09-2026",
            "driver_name": "K. Mohan Krishna",
            "total_km": "147",
            "starting_time": "08:00 AM",
            "closing_time": "05:30 PM",
            "extra_km": "67",
            "extra_hours": "1.5",
            "base_package": "8/80 = 2500",
            "extra_km_charge": "67 x 15 = 1005",
            "extra_hour_charge": "1.5 x 150 = 225",
            "bata": "250",
            "toll": "40",
            "total_amount": "4020",
        },
        "reviewer": "Human Reviewer",
        "notes": "Verified against physical scans with vehicle correction (6246 -> 6243)",
    }
    res_update = requests.put(
        f"{BASE_URL}/duty-slips/{DUTY_SLIP_ID}/extraction",
        json=review_payload
    )
    print(f"Status: {res_update.status_code}")
    assert res_update.status_code == 200, f"Update failed: {res_update.text}"
    updated_ext = res_update.json()
    print(f"Updated Status: {updated_ext['status']}")
    print(f"Audit Log Length: {len(updated_ext['audit_log'])}")
    print(f"Vehicle in Audit: Original={updated_ext['fields']['vehicle_number']['original_value']}, Current={updated_ext['fields']['vehicle_number']['value']}")

    # 4. Explicit Verify
    print("\n[4] POST Verify Extraction:")
    res_verify = requests.post(
        f"{BASE_URL}/duty-slips/{DUTY_SLIP_ID}/verify",
        json={"reviewer": "Human Reviewer", "notes": "Approved & Verified for Word Bill generation"}
    )
    print(f"Status: {res_verify.status_code}")
    assert res_verify.status_code == 200, f"Verification failed: {res_verify.text}"
    verified_ext = res_verify.json()
    print(f"Verified Status: {verified_ext['status']}")
    assert verified_ext["status"] == "verified"

    # 5. POST Generate Word Bill
    print("\n[5] POST Generate Word Bill:")
    res_gen = requests.post(f"{BASE_URL}/duty-slips/{DUTY_SLIP_ID}/generate-word")
    print(f"Status: {res_gen.status_code}")
    assert res_gen.status_code == 200, f"Word generation failed: {res_gen.text}"
    gen_data = res_gen.json()
    print(f"Generated Word Response:")
    print(f"  Filename: {gen_data['filename']}")
    print(f"  Storage Path: {gen_data['storage_path']}")
    print(f"  File Size: {gen_data['file_size_bytes']} bytes")
    print(f"  Duty Slip No: {gen_data['duty_slip_no']}")
    print(f"  Customer: {gen_data['customer_name']}")
    print(f"  Total Amount: {gen_data['total_amount']}")
    print(f"  Amount in words: {gen_data['amount_in_words']}")

    # 6. GET Download Word Bill
    print("\n[6] GET Download Generated .docx:")
    res_dl = requests.get(f"{BASE_URL}/duty-slips/{DUTY_SLIP_ID}/word-bill")
    print(f"Status: {res_dl.status_code}")
    assert res_dl.status_code == 200, f"Download failed: {res_dl.text}"
    docx_bytes = res_dl.content
    print(f"Downloaded DOCX bytes: {len(docx_bytes)}")

    # Save local copy for Word COM rendering
    local_docx = Path("./actual_generated_bill_47.docx").resolve()
    with open(local_docx, "wb") as f:
        f.write(docx_bytes)

    # 7. Inspect DOCX Structure
    doc = Document(str(local_docx))
    sec = doc.sections[0]
    print(f"\n[7] Document Inspection:")
    print(f"  Page Size: {sec.page_width.mm:.2f} mm x {sec.page_height.mm:.2f} mm (A4: 210 x 297)")
    print(f"  Tables Count: {len(doc.tables)}")
    for t_idx, table in enumerate(doc.tables):
        print(f"  Table {t_idx+1} ({len(table.rows)} rows x {len(table.columns)} columns):")
        for r_idx, row in enumerate(table.rows):
            print(f"    Row {r_idx}: {[c.text.strip().replace(chr(10), ' ') for c in row.cells]}")

    # 8. Render to PDF and PNG via Microsoft Word COM
    print("\n[8] Rendering Visual Screenshot via Microsoft Word COM...")
    local_pdf = local_docx.with_suffix(".pdf")
    word_app = win32com.client.Dispatch("Word.Application")
    word_app.Visible = False
    try:
        doc_com = word_app.Documents.Open(str(local_docx))
        # 17 = wdFormatPDF
        doc_com.SaveAs2(str(local_pdf), FileFormat=17)
        doc_com.Close()
    finally:
        word_app.Quit()

    pdf_doc = fitz.open(str(local_pdf))
    page = pdf_doc.load_page(0)
    pix = page.get_pixmap(dpi=300)

    artifact_dir = Path(r"C:\Users\uvish\.gemini\antigravity-ide\brain\f0390443-2b92-48de-994f-b1770d6b57b6")
    out_png_path = artifact_dir / "actual_live_word_bill_47.png"
    pix.save(str(out_png_path))
    print(f"Rendered Visual Screenshot saved at: {out_png_path}")

    # Clean up local temporary rendering files
    if local_docx.exists():
        local_docx.unlink()
    if local_pdf.exists():
        local_pdf.unlink()

    print("\n=== ACTUAL LIVE APPLICATION FLOW COMPLETED WITH 100% SUCCESS ===")


if __name__ == "__main__":
    run_actual_live_flow()
